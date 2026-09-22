"""Source regions and header context survive indexing, transport and an observed turn receipt."""
import asyncio
from io import BytesIO
from uuid import UUID

import pytest
from sqlalchemy import select, update
from docx import Document
from ax_workspace.entrypoints.mcp import McpReportsFacade
from ax_workspace.platform.persistence import ConversationContentEvidenceRecord, ConversationTurnRecord, MaterialChunkRecord, make_session_factory
from test_material_search import MINA, _stack, _upload

# 이 파일의 테스트는 **진짜 자식 프로세스**를 띄우고(자료·보고서 워커의 `IsolatedWork` spawn ·
# MCP `stdio_client` · `subprocess`) 그 진행을 초 단위 실시간 창으로 잰다 — 그래서 병렬 패스가 아니라
# `-n0` 직렬 패스에서 돈다. 기준과 걸개는 `tests/conftest.py`, 가르는 자리는 `Makefile` 의 `test-serial`.
pytestmark = pytest.mark.serial


@pytest.mark.parametrize("location", ["body", "header"])
def test_table_region_and_context_are_preserved_in_search_and_receipt(tmp_path, monkeypatch, location):
    client, application, worker, settings = _stack(tmp_path)
    task = client.post("/api/tasks", headers=MINA, json={"title": "표 근거"}).json()
    doc = Document()
    if location == "body":
        table = doc.add_table(rows=35, cols=2)
    else:
        from docx.shared import Inches
        table = doc.sections[0].header.add_table(rows=35, cols=2, width=Inches(4))
    table.cell(0, 0).text, table.cell(0, 1).text = "product", "deliverydeadline"
    for row in range(1, 35):
        table.cell(row, 0).text = f"row {row}"
        table.cell(row, 1).text = "scaxfinalrowtoken" if row == 34 else f"day {row}"
    stream = BytesIO()
    doc.save(stream)
    material = _upload(client, task["task_id"], "table.docx", stream.getvalue(), "application/octet-stream").json()
    assert asyncio.run(worker.run_once())
    query = "deliverydeadline scaxfinalrowtoken"
    found = client.get('/api/materials/search', headers=MINA, params={'q': query, 'limit': 1, 'resource_type': 'task', 'resource_id': task['task_id']}).json()
    hit = found["results"][0]
    assert hit["source_locator"]["row_start"] == 33 and hit["source_locator"]["row_end"] == 35
    assert hit["header_context"]["row"] == 1 and "deliverydeadline" in hit["header_context"]["text"]
    assert "scaxfinalrowtoken" in hit["excerpt"]
    if location == "header":
        assert hit["source_locator"]["section"] == 1
        assert hit["source_locator"]["container"] == "header"
    assert hit["extraction"]["coverage"]["complete"] is True
    conversation = client.post("/api/conversations", headers=MINA, json={"title": "표 위치"}).json()
    accepted = client.post(f"/api/conversations/{conversation['conversation_id']}/messages", headers={**MINA, "Idempotency-Key": "table-turn"}, json={"body": "표의 마지막 항목을 알려줘", "context": []}).json()
    with make_session_factory(settings.database_url)() as session:
        execution_id = session.get(ConversationTurnRecord, UUID(accepted["turn_id"])).execution_id
    monkeypatch.setenv("AX_MCP_CAUSATION_ID", str(execution_id))
    mcp = McpReportsFacade(settings, 'mina').search_materials(query, resource_type='task', resource_id=task['task_id'], limit=1)
    assert mcp == found
    view = client.get(f"/api/conversations/{conversation['conversation_id']}", headers=MINA).json()
    evidence = view["material_evidence"][0]
    assert evidence["source_locator"] == hit["source_locator"]
    assert evidence["header_context"] == hit["header_context"]
    assert evidence["extraction"] == hit["extraction"]
    resource = next(item for item in view["answer_resources"] if item["resource_id"] == material["material_id"])
    assert resource["source_locator"] == hit["source_locator"]

    # Rebuilding a stale index must retain header context as well as body terms.
    with make_session_factory(settings.database_url)() as session:
        session.execute(update(MaterialChunkRecord).where(
            MaterialChunkRecord.extraction_id == UUID(material["extraction"]["extraction_id"]),
        ).values(search_text="", analyzer_version="obsolete"))
        session.commit()
    assert application.reindex_material_search(limit=500) >= 3
    after = client.get('/api/materials/search', headers=MINA, params={'q': query, 'limit': 1, 'resource_type': 'task', 'resource_id': task['task_id']}).json()
    assert after == found

    # A purge must remove copied source text, including the newly stored header context.
    application.purge_attachment(application.authenticated_principal("yuna"), UUID(material["attachment_id"]), reason="합성 자료 삭제")
    with make_session_factory(settings.database_url)() as session:
        evidence_rows = list(session.scalars(select(ConversationContentEvidenceRecord).where(
            ConversationContentEvidenceRecord.attachment_id == UUID(material["attachment_id"]),
        )))
        assert evidence_rows and all(not row.excerpt and not row.header_context for row in evidence_rows)
