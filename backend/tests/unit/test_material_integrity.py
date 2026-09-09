"""A completed extraction must preserve searchable text at the end of the source."""
import pytest

from ax_workspace.platform.material_extraction import PypdfTextExtractor


@pytest.mark.parametrize("name, content_type", [("long.txt", "text/plain"), ("long.md", "text/markdown")])
def test_completed_text_keeps_content_beyond_the_old_character_limit(name, content_type):
    tail = "SCAX_UNIQUE_FINAL_PARAGRAPH_2026"
    source = "synthetic paragraph for full extraction.\n" * 6000 + tail
    outcome = PypdfTextExtractor().extract(name=name, content_type=content_type, data=source.encode())
    assert outcome.status == "completed"
    assert outcome.char_count == len(source)
    assert tail in "".join(block.text for block in outcome.blocks)
    assert any(tail in chunk.text for chunk in outcome.chunks)


def test_xlsx_keeps_late_rows_and_columns_with_original_coordinates():
    from io import BytesIO
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    for i in range(1, 2003):
        sheet.cell(i, 1, f"row {i}")
    sheet.cell(2002, 205, "SCAX_FINAL_CELL_205")
    stream = BytesIO()
    book.save(stream)
    outcome = PypdfTextExtractor().extract(name="wide.xlsx", content_type="", data=stream.getvalue())
    assert outcome.status == "completed"
    final = next(block for block in outcome.blocks if "SCAX_FINAL_CELL_205" in block.text)
    assert final.row == 2002 and final.sheet == "Sheet"
    assert "GW2002" in final.locator_label
    assert any("SCAX_FINAL_CELL_205" in chunk.text for chunk in outcome.chunks)


def test_docx_preserves_long_paragraphs_and_more_than_5000_blocks():
    from io import BytesIO
    from docx import Document

    doc = Document()
    doc.add_paragraph("word " * 1200 + "SCAX_LONG_PARAGRAPH_END")
    for i in range(5001):
        doc.add_paragraph(f"Synthetic distinct paragraph number {i:06d}")
    doc.add_paragraph("SCAX_FINAL_DOCX_PARAGRAPH")
    stream = BytesIO()
    doc.save(stream)
    outcome = PypdfTextExtractor().extract(name="long.docx", content_type="", data=stream.getvalue())
    assert outcome.status == "completed"
    assert any("SCAX_LONG_PARAGRAPH_END" in chunk.text for chunk in outcome.chunks)
    assert any("SCAX_FINAL_DOCX_PARAGRAPH" in chunk.text for chunk in outcome.chunks)


def test_pdf_reads_beyond_page_300_and_reports_unread_image_page():
    from test_document_parsers import _pdf

    source = _pdf(*[f"Synthetic page {i}" for i in range(300)], "SCAX_FINAL_PAGE", image_only_pages=1)
    outcome = PypdfTextExtractor().extract(name="long.pdf", content_type="", data=source)
    assert outcome.status == "partial"
    assert any(chunk.page == 301 and "SCAX_FINAL_PAGE" in chunk.text for chunk in outcome.chunks)
    assert outcome.coverage["missing_units"] == [{"page": 302, "reason": "needs_ocr"}]
    assert outcome.coverage["total_units"] == 302
    assert outcome.coverage["processed_units"] == 301
    assert outcome.warnings


def test_pptx_keeps_long_slide_table_and_final_notes():
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    box.text_frame.text = "synthetic slide " * 400 + "SCAX_SLIDE_END"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "synthetic cell " * 400 + "SCAX_TABLE_END"
    final = deck.slides.add_slide(deck.slide_layouts[6])
    final.notes_slide.notes_text_frame.text = "synthetic note " * 400 + "SCAX_NOTES_END"
    stream = BytesIO()
    deck.save(stream)
    outcome = PypdfTextExtractor().extract(name="deck.pptx", content_type="", data=stream.getvalue())
    try:
        assert outcome.status == "completed"
        for token in ("SCAX_SLIDE_END", "SCAX_TABLE_END", "SCAX_NOTES_END"):
            assert any(token in chunk.text for chunk in outcome.chunks)
        assert outcome.blocks[-1].slide == 2 and outcome.blocks[-1].kind == "notes"
    finally:
        outcome.close()


def test_custom_output_budget_is_never_published_as_complete():
    from ax_workspace.modules.work.document_parsing import ParseBudget
    from ax_workspace.platform.document_parsers import OfficeDocumentParser
    from test_document_parsers import _docx

    outcome = PypdfTextExtractor(OfficeDocumentParser(ParseBudget(max_blocks=1))).extract(
        name="bounded.docx", content_type="", data=_docx())
    assert outcome.status == "too_large" and not outcome.blocks and not outcome.chunks
    assert outcome.coverage["complete"] is False and outcome.warnings


def test_docx_keeps_nested_tables_and_header_tables():
    from io import BytesIO
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "SCAX_NESTED_TABLE"
    doc.sections[0].header.add_table(rows=1, cols=1, width=Inches(3)).cell(0, 0).text = "SCAX_HEADER_TABLE"
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.paragraphs[0].text = "SCAX_FIRST_PAGE_HEADER"
    stream = BytesIO()
    doc.save(stream)
    outcome = PypdfTextExtractor().extract(name="nested.docx", content_type="", data=stream.getvalue())
    try:
        assert outcome.status == "completed"
        for token in ("SCAX_NESTED_TABLE", "SCAX_HEADER_TABLE", "SCAX_FIRST_PAGE_HEADER"):
            assert any(token in chunk.text for chunk in outcome.chunks)
    finally:
        outcome.close()


@pytest.mark.parametrize("extension", ["docx", "xlsx", "pptx", "pdf", "txt"])
def test_temporary_spool_failure_retries_instead_of_marking_source_corrupt(extension, monkeypatch):
    from uuid import uuid4
    from test_document_parsers import _docx, _xlsx, _pptx, _pdf
    from ax_workspace.modules.work.material_extraction import ClaimedExtraction, MaterialExtractionService
    from ax_workspace.platform.extraction_spool import ExtractionSpool

    data = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx, "pdf": lambda: _pdf("Readable"), "txt": lambda: b"Readable"}[extension]()

    class Storage:
        def get(self, ref):
            return data

    def disk_full(self, record):
        raise OSError("synthetic disk full")

    monkeypatch.setattr(ExtractionSpool, "append", disk_full)
    claimed = ClaimedExtraction(uuid4(), 1, f"source.{extension}", "", "synthetic")
    outcome = MaterialExtractionService.extract(claimed, Storage(), PypdfTextExtractor())
    assert outcome.transient and outcome.failure_reason == "extractor_error"


def test_pdf_resource_limit_is_too_large_not_recoverable_partial(monkeypatch):
    from pypdf import PageObject
    from pypdf.errors import LimitReachedError
    from test_document_parsers import _pdf

    calls = 0
    def extract(self, *args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise LimitReachedError("synthetic decompression ceiling")
        return "readable first page"
    monkeypatch.setattr(PageObject, "extract_text", extract)
    outcome = PypdfTextExtractor().extract(name="limit.pdf", content_type="", data=_pdf("First", "Second"))
    assert outcome.status == "too_large" and not outcome.chunks and not outcome.blocks
    assert outcome.failure_reason == "too_large"


def test_pdf_admission_checks_decoded_content_before_text_parser(monkeypatch):
    from pypdf import PageObject
    from ax_workspace.modules.work.document_parsing import ParseBudget
    from ax_workspace.platform.document_parsers import OfficeDocumentParser
    from test_document_parsers import _pdf

    def must_not_parse(self, *args, **kwargs):
        raise AssertionError("oversize content must be refused before text extraction")
    monkeypatch.setattr(PageObject, "extract_text", must_not_parse)
    outcome = PypdfTextExtractor(OfficeDocumentParser(ParseBudget(max_pdf_content_bytes=100))).extract(
        name="large-stream.pdf", content_type="", data=_pdf("synthetic " * 100))
    assert outcome.status == "too_large" and not outcome.chunks
    assert outcome.coverage["complete"] is False


def test_pdf_nested_form_is_included_in_admission_before_extract(monkeypatch):
    from io import BytesIO
    from pypdf import PdfWriter, PageObject
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    from ax_workspace.modules.work.document_parsing import ParseBudget
    from ax_workspace.platform.document_parsers import OfficeDocumentParser

    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    form = DecodedStreamObject()
    form.set_data(b"synthetic form content " * 20)
    form[NameObject("/Subtype")] = NameObject("/Form")
    reference = writer._add_object(form)
    resources = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Fm"): reference})})
    # The cycle must not make admission loop or omit the form's own decoded bytes.
    form[NameObject("/Resources")] = resources
    page[NameObject("/Resources")] = resources
    stream = BytesIO()
    writer.write(stream)
    calls = []
    monkeypatch.setattr(PageObject, "extract_text", lambda *args, **kwargs: calls.append(1) or "should not parse")
    outcome = PypdfTextExtractor(OfficeDocumentParser(ParseBudget(max_pdf_content_bytes=100))).extract(
        name="form.pdf", content_type="", data=stream.getvalue())
    assert outcome.status == "too_large" and calls == []


def test_pdf_document_admission_drops_preceding_pages_instead_of_publishing_partial():
    from ax_workspace.modules.work.document_parsing import ParseBudget
    from ax_workspace.platform.document_parsers import OfficeDocumentParser
    from test_document_parsers import _pdf

    outcome = PypdfTextExtractor(OfficeDocumentParser(ParseBudget(
        max_pdf_content_bytes=1000, max_pdf_total_content_bytes=200,
    ))).extract(name="total.pdf", content_type="", data=_pdf("word " * 20, "word " * 20))
    assert outcome.status == "too_large" and not outcome.blocks and not outcome.chunks
    assert outcome.coverage["complete"] is False


def _form_pdf(repeats=1):
    from io import BytesIO
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    fonts = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    def form(text):
        item = DecodedStreamObject()
        item.set_data(f"BT /F1 12 Tf ({text}) Tj ET".encode())
        item[NameObject("/Subtype")] = NameObject("/Form")
        item[NameObject("/Resources")] = fonts
        return writer._add_object(item)
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Fm"): form("Repeated"), NameObject("/Tail"): form("SCAX_TAIL_FORM_TOKEN")})})
    content = DecodedStreamObject()
    content.set_data(b"/Fm Do\n" * repeats + b"/Tail Do\n")
    page[NameObject("/Contents")] = writer._add_object(content)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


@pytest.mark.parametrize("repeats", [1, 5000])
@pytest.mark.parametrize("pdf_log_level", [30, 50], indirect=True)
def test_pdf_form_invocation_ceiling_cannot_hide_a_final_form(repeats, pdf_log_level):
    outcome = PypdfTextExtractor().extract(name="forms.pdf", content_type="", data=_form_pdf(repeats))
    try:
        if repeats == 1:
            assert outcome.status == "completed" and any("SCAX_TAIL_FORM_TOKEN" in chunk.text for chunk in outcome.chunks)
        else:
            assert outcome.status == "too_large" and not outcome.chunks and not outcome.blocks
            assert outcome.coverage["complete"] is False
    finally:
        outcome.close()


def test_pdf_warnings_do_not_contaminate_another_workers_document(monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    import logging
    from threading import Barrier
    from pypdf import PageObject
    from test_document_parsers import _pdf

    barrier = Barrier(2)
    original = PageObject.extract_text
    def extract(page, *args, **kwargs):
        text = original(page, *args, **kwargs)
        barrier.wait(timeout=5)
        if "noisy" in text:
            logging.getLogger("pypdf._page").warning("synthetic undecodable source fragment")
        barrier.wait(timeout=5)
        return text
    monkeypatch.setattr(PageObject, "extract_text", extract)
    def run(label):
        return PypdfTextExtractor().extract(name="thread.pdf", content_type="", data=_pdf(label))
    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(run, ["clean", "noisy"]))
    try:
        assert outcomes[0].status == "completed" and not outcomes[0].warnings
        assert outcomes[1].status == "partial" and outcomes[1].coverage["missing_units"] == [{"page": 1, "reason": "parser_warning"}]
    finally:
        for outcome in outcomes:
            outcome.close()


@pytest.mark.parametrize("pdf_log_level", [30, 50], indirect=True)
def test_pdf_capacity_failure_inside_form_is_not_recoverable_partial(monkeypatch, pdf_log_level):
    import logging
    from pypdf import PageObject
    from pypdf.errors import LimitReachedError

    original = PageObject.extract_xform_text
    calls = 0
    def fail_form(self, *args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            return original(self, *args, **kwargs)
        raise LimitReachedError("synthetic nested decoder budget")
    monkeypatch.setattr(PageObject, "extract_xform_text", fail_form)
    outcome = PypdfTextExtractor().extract(name="nested-limit.pdf", content_type="", data=_form_pdf())
    try:
        assert outcome.status == "too_large" and not outcome.chunks and not outcome.blocks
    finally:
        outcome.close()


@pytest.fixture
def pdf_log_level(request):
    import logging
    logger = logging.getLogger("pypdf._page")
    previous = logger.level
    logger.setLevel(request.param)
    try:
        yield request.param
    finally:
        logger.setLevel(previous)
