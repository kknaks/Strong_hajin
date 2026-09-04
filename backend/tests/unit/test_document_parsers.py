"""Office/PDF parser: ordered blocks with source locators, explainable statuses, and hard resource budgets.

All fixtures are synthetic; no real company document is read here.
"""
from io import BytesIO
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from ax_workspace.modules.work.document_parsing import (
    STATUS_BUDGET_EXCEEDED,
    STATUS_CORRUPT,
    STATUS_EMPTY,
    STATUS_ENCRYPTED,
    STATUS_NEEDS_OCR,
    STATUS_OK,
    STATUS_UNSUPPORTED,
    ParseBudget,
    classify_document,
)
from ax_workspace.platform.document_parsers import OfficeDocumentParser

PARSER = OfficeDocumentParser()


def _docx(*, paragraphs_then_table=True) -> bytes:
    document = Document()
    document.add_paragraph("첫 문단: 계약 목적")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "항목", "금액"
    table.cell(1, 0).text, table.cell(1, 1).text = "대행료", "1,000,000"
    document.add_paragraph("둘째 문단: 지급 조건")
    document.sections[0].header.paragraphs[0].text = "머리글 문구"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "매출"
    sheet["A1"], sheet["B1"], sheet["C1"] = "일자", "건수", "합계"
    sheet["A2"], sheet["B2"], sheet["C2"] = "2026-09-01", 3, "=SUM(B2:B2)*1000"
    sheet.append([])  # blank row is skipped, not a block
    sheet["A4"] = "비고: 부가세 별도"
    hidden = workbook.create_sheet("숨김")
    hidden["A1"] = "숨겨진 값"
    hidden.sheet_state = "hidden"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "체험단 보고"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "본문 텍스트"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "채널", "건수"
    table.cell(1, 0).text, table.cell(1, 1).text = "블로그", "12"
    slide.notes_slide.notes_text_frame.text = "발표자 노트 내용"
    presentation.slides.add_slide(presentation.slide_layouts[6])  # empty slide contributes nothing
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf(*pages: str, image_only_pages: int = 0) -> bytes:
    """Minimal PDF: text pages via a Helvetica stream, image-only pages via a 1x1 XObject and no text."""
    objects: list[str] = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids: list[int] = []
    image_obj = None
    if image_only_pages:
        objects.append("<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\nstream\n\x00\nendstream")
        image_obj = len(objects)
    for text in pages:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
        content_obj = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {content_obj} 0 R >>")
        kids.append(len(objects))
    for _ in range(image_only_pages):
        content = "q 100 0 0 100 50 50 cm /Im1 Do Q"
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
        content_obj = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /XObject << /Im1 {image_obj} 0 R >> >> /Contents {content_obj} 0 R >>")
        kids.append(len(objects))
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out = "%PDF-1.4\n"
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out.encode("latin-1")))
        out += f"{index} 0 obj\n{body}\nendobj\n"
    xref = len(out.encode("latin-1"))
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n" + "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


def test_classification_covers_office_and_pdf_only() -> None:
    assert classify_document("계약서.docx", "") == "docx" and classify_document("x.xlsm", "") == "xlsx"
    assert classify_document("deck.pptx", "") == "pptx" and classify_document("scan.PDF", "") == "pdf"
    assert classify_document("blob", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") == "xlsx"
    assert classify_document("notes.md", "text/markdown") is None and classify_document("old.doc", "") is None


def test_docx_keeps_paragraph_and_table_order_and_headers() -> None:
    result = PARSER.parse(name="계약.docx", content_type="", data=_docx())
    assert result.status == STATUS_OK
    assert [block.kind for block in result.blocks] == ["paragraph", "table", "paragraph", "header"]
    assert [block.locator.label for block in result.blocks] == ["문단 1", "표 1", "문단 2", "머리글 1"]
    assert result.blocks[1].text == "항목 | 금액 대행료 | 1,000,000"
    assert [block.order for block in result.blocks] == [0, 1, 2, 3]


def test_xlsx_returns_sheet_and_cell_ranges_and_keeps_formulas_as_text() -> None:
    result = PARSER.parse(name="매출.xlsx", content_type="", data=_xlsx())
    assert result.status == STATUS_OK and result.sheet_count == 2
    first = result.blocks[0]
    assert first.locator.kind == "sheet_row" and first.locator.sheet == "매출" and first.locator.cell_range == "A1:C1"
    assert first.text == "A1=일자 | B1=건수 | C1=합계"
    formula_row = result.blocks[1]
    assert "C2==SUM(B2:B2)*1000" in formula_row.text  # the formula string, never a computed number
    assert result.blocks[2].locator.cell_range == "A4:A4"
    assert any(block.locator.sheet == "숨김" for block in result.blocks)
    assert "hidden sheet included: 숨김" in result.warnings and any("formulas kept as text" in w for w in result.warnings)


def test_pptx_returns_slide_numbers_tables_and_notes() -> None:
    result = PARSER.parse(name="보고.pptx", content_type="", data=_pptx())
    assert result.status == STATUS_OK and result.slide_count == 2
    kinds = [(block.kind, block.locator.label) for block in result.blocks]
    assert kinds == [("slide_text", "슬라이드 1"), ("slide_table", "슬라이드 1"), ("notes", "슬라이드 1 발표자 노트")]
    assert "체험단 보고" in result.blocks[0].text and "본문 텍스트" in result.blocks[0].text
    assert result.blocks[1].text == "채널 | 건수 블로그 | 12"
    assert result.blocks[2].text == "발표자 노트 내용"


def test_pdf_pages_needs_ocr_and_empty_are_distinct() -> None:
    text = PARSER.parse(name="a.pdf", content_type="application/pdf", data=_pdf("Page one", "Page two"))
    assert text.status == STATUS_OK and [block.locator.label for block in text.blocks] == ["1쪽", "2쪽"]
    scan = PARSER.parse(name="scan.pdf", content_type="application/pdf", data=_pdf(image_only_pages=2))
    assert scan.status == STATUS_NEEDS_OCR and scan.blocks == () and scan.page_count == 2 and "OCR" in (scan.detail or "")
    blank = PARSER.parse(name="blank.pdf", content_type="application/pdf", data=_pdf(""))
    assert blank.status == STATUS_EMPTY and blank.blocks == ()
    mixed = PARSER.parse(name="mixed.pdf", content_type="application/pdf", data=_pdf("Cover text", image_only_pages=1))
    assert mixed.status == STATUS_OK and len(mixed.blocks) == 1 and any("no text layer" in w for w in mixed.warnings)


def test_corrupt_encrypted_and_legacy_inputs_end_in_explainable_statuses() -> None:
    assert PARSER.parse(name="x.docx", content_type="", data=b"PK\x03\x04 not really a zip").status == STATUS_CORRUPT
    assert PARSER.parse(name="x.docx", content_type="", data=b"plain text").status == STATUS_CORRUPT
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    legacy = PARSER.parse(name="old.xlsx", content_type="", data=ole)
    assert legacy.status == STATUS_UNSUPPORTED and "legacy" in (legacy.detail or "")
    encrypted = PARSER.parse(name="locked.docx", content_type="", data=ole + "EncryptedPackage".encode("utf-16-le"))
    assert encrypted.status == STATUS_ENCRYPTED
    # A valid zip whose parts are not a document opens but fails as corrupt, not as an exception.
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<broken")
    assert PARSER.parse(name="broken.docx", content_type="", data=buffer.getvalue()).status == STATUS_CORRUPT
    assert PARSER.parse(name="image.png", content_type="image/png", data=b"\x89PNG").status == STATUS_UNSUPPORTED


def test_ooxml_preflight_refuses_bombs_before_opening() -> None:
    tight = OfficeDocumentParser(ParseBudget(max_member_bytes=10_000, max_total_uncompressed_bytes=20_000, max_archive_members=3))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 50_000)  # 50KB of zeros compresses to a few hundred bytes
    bomb = tight.parse(name="bomb.docx", content_type="", data=buffer.getvalue())
    assert bomb.status == STATUS_BUDGET_EXCEEDED and bomb.blocks == () and "exceeds" in (bomb.detail or "")
    many = BytesIO()
    with zipfile.ZipFile(many, "w") as archive:
        for index in range(5):
            archive.writestr(f"part{index}.xml", "x")
    assert tight.parse(name="many.docx", content_type="", data=many.getvalue()).status == STATUS_BUDGET_EXCEEDED
    # Encrypted zip members are refused as encrypted, not decompressed.
    enc = BytesIO()
    with zipfile.ZipFile(enc, "w") as archive:
        archive.writestr("word/document.xml", "x")
    raw = bytearray(enc.getvalue())
    # zipfile does not write the encryption flag itself; set bit 0 of the general-purpose flags in both headers.
    local = raw.index(b"PK\x03\x04")
    raw[local + 6] |= 0x1
    central = raw.index(b"PK\x01\x02")
    raw[central + 8] |= 0x1
    assert tight.parse(name="enc.docx", content_type="", data=bytes(raw)).status == STATUS_ENCRYPTED


def test_output_budget_truncates_with_a_warning_instead_of_growing_unbounded() -> None:
    small = OfficeDocumentParser(ParseBudget(max_blocks=2, max_chars=10_000))
    result = small.parse(name="계약.docx", content_type="", data=_docx())
    assert result.status == STATUS_OK and len(result.blocks) == 2 and result.truncated
    assert "output truncated at budget" in result.warnings
    rows = OfficeDocumentParser(ParseBudget(max_rows_per_sheet=1))
    sheet = rows.parse(name="매출.xlsx", content_type="", data=_xlsx())
    assert any("sheet rows truncated at 1" in w for w in sheet.warnings)
