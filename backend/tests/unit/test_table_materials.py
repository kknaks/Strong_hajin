"""Table excerpts retain exact source regions and explicitly identified context rows."""
from io import BytesIO

import pytest

from ax_workspace.platform.material_extraction import PypdfTextExtractor


def wide_sheet():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    sheet.append([f"header_{column}" for column in range(1, 66)])
    sheet.append([f"value_{column}" for column in range(1, 66)])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_wide_sheet_preserves_disjoint_column_regions_and_their_context_rows():
    outcome = PypdfTextExtractor().extract(name="wide.xlsx", content_type="", data=wide_sheet())
    try:
        rows = [block for block in outcome.blocks if block.row == 2]
        assert len(rows) == 3
        assert [(block.source_locator["column_start"], block.source_locator["column_end"]) for block in rows] == [(1, 32), (33, 64), (65, 65)]
        assert [block.source_locator["cell_range"] for block in rows] == ["A2:AF2", "AG2:BL2", "BM2:BM2"]
        assert "header_65" in rows[-1].header_context["text"]
        assert rows[-1].header_context["basis"] == "first_non_empty_row"
        assert rows[-1].header_context["row"] == 1
        assert "value_65" in rows[-1].text
    finally:
        outcome.close()


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_long_tables_keep_row_regions_and_first_row_context(kind):
    if kind == "docx":
        from docx import Document
        document = Document()
        table = document.add_table(rows=35, cols=2)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        table = slide.shapes.add_table(35, 2, Inches(1), Inches(1), Inches(4), Inches(4)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "product", "deliverydeadline"
    for row in range(1, 35):
        table.cell(row, 0).text = f"item {row}"
        table.cell(row, 1).text = "scaxfinalrowtoken" if row == 34 else f"day {row}"
    stream = BytesIO()
    document.save(stream)
    outcome = PypdfTextExtractor().extract(name=f"table.{kind}", content_type="", data=stream.getvalue())
    try:
        blocks = [block for block in outcome.blocks if block.kind in {"table", "slide_table"}]
        assert [(b.source_locator["row_start"], b.source_locator["row_end"]) for b in blocks] == [(1, 16), (17, 32), (33, 35)]
        assert all("deliverydeadline" in b.header_context["text"] for b in blocks)
        assert all(b.source_locator["index"] == 1 for b in blocks)
        assert "scaxfinalrowtoken" in blocks[-1].text
        if kind == "pptx":
            assert blocks[-1].source_locator["slide"] == 1
    finally:
        outcome.close()


def test_docx_merged_context_spans_column_batches_without_moving_omitted_cells():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    table = doc.add_table(rows=2, cols=35)
    table.cell(0, 0).merge(table.cell(0, 34)).text = "groupheading"
    # The second row omits column 1; its first actual cell is column 2, not a shifted column 1.
    row = table.rows[1]
    row._tr.remove(row._tr.tc_lst[0])
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    row._tr.get_or_add_trPr().append(before)
    row.cells[0].text = "firstactualcell"
    row.cells[-1].text = "lastactualcell"
    stream = BytesIO()
    doc.save(stream)
    outcome = PypdfTextExtractor().extract(name="merged.docx", content_type="", data=stream.getvalue())
    try:
        last = next(block for block in outcome.blocks if "lastactualcell" in block.text)
        assert last.source_locator["column_start"] == 33 and last.source_locator["column_end"] == 35
        assert "groupheading" in last.header_context["text"]
        assert any("R2C2=firstactualcell" in block.text for block in outcome.blocks)
    finally:
        outcome.close()


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_sparse_table_values_do_not_lose_which_column_they_belong_to(kind):
    def extract(column):
        if kind == "docx":
            from docx import Document
            doc = Document()
            table = doc.add_table(rows=2, cols=3)
        else:
            from pptx import Presentation
            from pptx.util import Inches
            doc = Presentation()
            slide = doc.slides.add_slide(doc.slide_layouts[6])
            table = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(3), Inches(2)).table
        for col in range(3):
            table.cell(0, col).text = f"header{col + 1}"
        table.cell(1, column - 1).text = "Value"
        stream = BytesIO()
        doc.save(stream)
        outcome = PypdfTextExtractor().extract(name=f"sparse.{kind}", content_type="", data=stream.getvalue())
        try:
            return [(b.text, b.source_locator, b.header_context) for b in outcome.blocks]
        finally:
            outcome.close()
    assert extract(2) != extract(3)


def test_xlsx_merged_header_reaches_the_final_column_region():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.merge_cells("A1:BM1")
    sheet["A1"] = "sharedheading"
    sheet["BM2"] = "finalcolumn"
    stream = BytesIO()
    book.save(stream)
    outcome = PypdfTextExtractor().extract(name="merged.xlsx", content_type="", data=stream.getvalue())
    try:
        last = next(b for b in outcome.blocks if "finalcolumn" in b.text)
        assert "sharedheading" in last.header_context["text"]
        assert last.source_locator["cell_range"] == "BM2:BM2"
    finally:
        outcome.close()
