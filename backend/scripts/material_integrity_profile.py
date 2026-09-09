"""Synthetic-only extraction profile. Run each kind in a fresh process; no external corpus is read.

uv run python scripts/material_integrity_profile.py xlsx
"""
from __future__ import annotations

import argparse
import asyncio
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import resource
import sys
from tempfile import TemporaryDirectory
from time import perf_counter
from zipfile import ZipFile

from fastapi.testclient import TestClient
from sqlalchemy import func, select

from ax_workspace.bootstrap.material_worker import MaterialExtractionWorker
from ax_workspace.bootstrap.settings import RuntimeProfile, Settings
from ax_workspace.entrypoints.http import create_app
from ax_workspace.entrypoints.reset_demo import reset_database
from ax_workspace.platform.material_extraction import PypdfTextExtractor
from ax_workspace.platform.persistence import MaterialBlockRecord, MaterialChunkRecord, make_session_factory

TAIL = "scaxuniquefinaltoken"


def fixture(kind: str, slides: int = 1000) -> tuple[bytes, dict]:
    stream = BytesIO()
    if kind == "xlsx":
        from openpyxl import Workbook
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Synthetic")
        sheet.append([f"column{i}" for i in range(10)])
        for row in range(1, 100_000):
            values = [row * 10 + col for col in range(10)]
            if row == 99_999:
                values[-1] = TAIL
            sheet.append(values)
        book.save(stream)
        shape = {"rows": 100_000, "non_empty_cells": 1_000_000}
    elif kind == "docx":
        from docx import Document
        doc = Document()
        for page in range(350):
            doc.add_paragraph(f"Synthetic page {page} " + sha256(str(page).encode()).hexdigest() * 10)
            table = doc.add_table(rows=12, cols=4)
            for row in range(12):
                for col in range(4):
                    table.cell(row, col).text = f"cell {page}-{row}-{col}"
            if page < 349:
                doc.add_page_break()
        doc.add_paragraph(TAIL)
        doc.save(stream)
        shape = {"authored_pages": 350, "tables": 350, "non_empty_cells": 16_800}
    else:
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        for number in range(slides):
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = f"Synthetic slide {number}"
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(2)).table
            table.cell(1, 1).text = f"Synthetic cell {number}"
            slide.notes_slide.notes_text_frame.text = TAIL if number == slides - 1 else f"Synthetic notes {number}"
        deck.save(stream)
        shape = {"slides": slides, "tables": slides, "notes": slides}
    return stream.getvalue(), shape


def main() -> None:
    args = argparse.ArgumentParser()
    args.add_argument("kind", choices=["xlsx", "docx", "pptx"])
    args.add_argument("--slides", type=int, default=1000)
    options = args.parse_args()
    kind = options.kind
    if options.slides < 1:
        args.error("--slides must be positive")
    began = perf_counter()
    data, shape = fixture(kind, slides=options.slides)
    generation_seconds = perf_counter() - began
    with ZipFile(BytesIO(data)) as package:
        infos = package.infolist()
        archive = {"members": len(infos), "uncompressed_bytes": sum(i.file_size for i in infos),
                   "largest_member_bytes": max(i.file_size for i in infos)}
    result = {"profile": "scax-material-integrity-v1", "kind": kind, "shape": shape,
              "source_bytes": len(data), "sha256": sha256(data).hexdigest(), "archive": archive,
              "generation_seconds": round(generation_seconds, 3)}
    print(json.dumps({"phase": "generated", "kind": kind, "source_bytes": len(data)}), file=sys.stderr, flush=True)
    with TemporaryDirectory(prefix="scax-material-profile-") as root:
        directory = Path(root)
        url = f"sqlite:///{directory / 'test.db'}"
        reset_database(url)
        settings = Settings(RuntimeProfile.TEST, url, materials_dir=str(directory / "materials"))
        app = create_app(settings)
        client = TestClient(app)
        headers = {"X-Demo-Persona": "mina"}
        task = client.post("/api/tasks", headers=headers, json={"title": "Synthetic integrity profile"}).json()
        uploaded = client.post(f"/api/tasks/{task['task_id']}/materials", headers=headers, data={"kind": "input"},
                               files={"file": (f"synthetic.{kind}", data, "application/octet-stream")})
        assert uploaded.status_code == 201, uploaded.status_code
        del data

        class TimedExtractor:
            def extract(self, **kwargs):
                start = perf_counter()
                outcome = PypdfTextExtractor().extract(**kwargs)
                result["parse_seconds"] = round(perf_counter() - start, 3)
                result["parser_status"] = outcome.status
                print(json.dumps({"phase": "parsed", "kind": kind, "seconds": result["parse_seconds"], "status": outcome.status}), file=sys.stderr, flush=True)
                return outcome

        worker = MaterialExtractionWorker(settings, extractor=TimedExtractor(), queue_factory=lambda session: app.state.workflow_application.memory_job_queue)
        start = perf_counter()
        assert asyncio.run(worker.run_once())
        result["worker_seconds"] = round(perf_counter() - start, 3)
        extraction = client.get(f"/api/tasks/{task['task_id']}/materials", headers=headers).json()[0]["extraction"]
        result["extraction"] = {k: extraction[k] for k in ("status", "failure_reason", "char_count", "chunk_count", "coverage", "warnings")}
        with make_session_factory(url)() as session:
            result["db_blocks"] = session.scalar(select(func.count()).select_from(MaterialBlockRecord))
            result["db_chunks"] = session.scalar(select(func.count()).select_from(MaterialChunkRecord))
        search_start = perf_counter()
        search = client.get('/api/materials/search', headers=headers, params={'q': TAIL, 'resource_type': 'task', 'resource_id': task['task_id']}).json()
        result["tail_search_seconds"] = round(perf_counter() - search_start, 3)
        result["tail_found"] = any(TAIL in hit["excerpt"] for hit in search["results"])
        if extraction["status"] == "completed":
            assert result["tail_found"], "completed projection lost tail token"
    result["peak_rss_bytes_including_generation"] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * (1 if sys.platform == "darwin" else 1024)
    result["total_seconds"] = round(perf_counter() - began, 3)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
